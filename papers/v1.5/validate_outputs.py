from pathlib import Path
import re
import docx

BASE_DIR = Path(__file__).resolve().parent
DOCX_PATH = BASE_DIR / 'v1.5_KSII_TIIS.docx'
MD_PATH = BASE_DIR / 'v1.5.md'

print(f"=== Checking {DOCX_PATH.name} ===")
doc = docx.Document(str(DOCX_PATH))
print('Paragraphs in KO docx:', len(doc.paragraphs))
print('Tables in KO docx:', len(doc.tables))
for i in [0, 8, 11, 42, 43, 74, 94, 100, 106, 116, 122]:
    p = doc.paragraphs[i]
    has_g = '<a:graphic' in p._p.xml or 'w:drawing' in p._p.xml
    print(f'P{i}: style={p.style.name}, text={p.text[:60]}, has_graphic={has_g}')

for t_idx, table in enumerate(doc.tables):
    print(f'Table {t_idx}: {len(table.rows)} rows, col0={table.rows[0].cells[0].text[:30]}')

print(f"\n=== Checking {MD_PATH.name} ===")
md_text = MD_PATH.read_text(encoding='utf-8')

unicode_quotes = ['\u201c', '\u201d', '\u2018', '\u2019']
found_uq = [uq for uq in unicode_quotes if uq in md_text]
print('Unicode quotes found:', found_uq)

bold_matches = re.findall(r'\*\*.*?\*\*', md_text)
print('Bold ** matches:', len(bold_matches), bold_matches[:5])
under_bold = re.findall(r'__.*?__', md_text)
print('Under bold __ matches:', len(under_bold), under_bold[:5])

italic_matches = re.findall(r'(?<!\*)\*([^*\s][^*]*[^*\s])\*(?!\*)', md_text)
print('Potential italic * matches:', len(italic_matches), italic_matches[:5])

print('Total lines in MD:', len(md_text.splitlines()))
print('Validation complete.')
