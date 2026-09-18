"""Populate the user's two-page Word template only after the main result audit."""
from copy import deepcopy
from pathlib import Path
import shutil
from collections import Counter
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

from rttdist.experiment_io import read_json, write_json, digest, timestamp

ROOT=Path(__file__).resolve().parents[1]
LABEL={'cpp':'C','haskell':'H','prolog':'P'}

def font(run,size=10,bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.bold=bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'바탕')

def paragraph(doc,text,size=10,bold=False,after=4,keep=False):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    f=p.paragraph_format;f.space_after=Pt(after);f.space_before=Pt(0);f.line_spacing=1
    f.keep_with_next=keep;f.widow_control=True
    font(p.add_run(text),size,bold)
    return p

def main():
    root=ROOT/'artifacts-lmstudio/fps-v1/main'
    summary=read_json(root/'summary.json');audit=read_json(root/'completion_audit.json')
    assert summary['complete'] and audit['passed'] and audit['observed']==90
    records=[read_json(p) for p in sorted(root.glob('*/*/observation.json'))]
    ok=[r for r in records if r['success']]
    counts=Counter(r['status'] for r in records)
    destination=ROOT/'docs/abs/33rd_abstract_sample.docx'
    backup=ROOT/'docs/abs/.source-backup/33rd_abstract_before_results.docx'
    backup.parent.mkdir(exist_ok=True)
    if not backup.exists():shutil.copy2(destination,backup)
    doc=Document(backup)
    first_section=deepcopy(doc.sections[0]._sectPr)
    last_section=deepcopy(doc.sections[-1]._sectPr)
    for child in list(doc._element.body):doc._element.body.remove(child)
    doc._element.body.append(last_section)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10)
    normal.paragraph_format.line_spacing=1
    paragraph(doc,'왕복 번역의 기능 보존과 표현 안정화:\n세 프로그래밍 언어의 FPS 거리 분석',20,True,8,True)
    distances=[r['d_n'] for r in ok]
    dtext=f'{min(distances)}~{max(distances)}' if distances else '미관측'
    paragraph(doc,
        f'(Abstract) 코드 왕복 번역에서 기능 보존과 표현의 안정화를 분리하여 측정하였다. '
        f'C++·Haskell·Prolog의 공통 알고리즘 문제 15개를 대상으로 Qwen2.5-Coder 7B를 사용해 '
        f'6방향, 총 90건의 본 실험을 수행했다. 각 생성 코드에 13개 평가 입력을 적용하고, '
        f'인접 복원 코드의 JPlag text 유사도가 0.85 이상이면 근사 고정점으로 판정했다. '
        f'거리는 출발 코드를 포함하여 관측한 서로 다른 언어·소스 상태의 수로 정의했다. '
        f'{len(ok)}건이 기능 검사와 수렴 조건을 함께 통과했으며 성공 거리의 범위는 {dtext}였다. '
        '실패의 조기 종료를 짧은 거리로 처리하지 않고, 성공률·성공 조건부 거리·원본 대비 유사도를 함께 보고한다. '
        '결과는 이 모델과 초기 구현에서의 번역 안정성을 나타내며 언어 자체의 보편적 거리로 일반화하지 않는다.',10,True,7)
    boundary=doc.add_paragraph();boundary.paragraph_format.space_after=Pt(0)
    boundary._p.get_or_add_pPr().append(first_section)
    paragraph(doc,'1. 서론',11,True,4,True)
    paragraph(doc,'코드 번역은 문법뿐 아니라 자료 구조, 제어 흐름과 입출력 동작의 보존을 요구한다. '
        '자연어 왕복 번역과 고정점 집합(Fixed-Point Set, FPS)의 크기를 이용한 선행 연구[1]를 바탕으로, '
        '실행 가능한 코드의 기능 검사와 근사 수렴을 결합하였다. 번역 실패가 빨리 발생한 경우와 '
        '기능을 유지하면서 적은 상태로 안정화한 경우를 구분하는 것이 본 측정의 목적이다.')
    paragraph(doc,'2. 거리와 실험 방법',11,True,4,True)
    paragraph(doc,'출발 언어 A의 원본을 x₀, t번째 중간 언어 B의 코드를 yₜ, A로 복원한 코드를 xₜ라 두면 '
        '경로는 x₀ → y₁ → x₁ → y₂ → x₂ → …이다. 편도 j단계에서 생성한 언어·소스를 (Lⱼ,zⱼ)라 할 때, '
        '줄바꿈만 LF로 통일한 소스의 SHA-256 해시 h로 상태 집합을 정의한다.')
    paragraph(doc,'S₀ = {(A,h(x₀))}\nSⱼ = Sⱼ₋₁ ∪ {(Lⱼ,h(zⱼ))},  Dⱼ = |Sⱼ|',10,False,5)
    paragraph(doc,'주석·공백·식별자가 다르면 별개 상태이며, 같은 소스도 언어가 다르면 구분한다. '
        '유사도 임계값은 상태 중복 제거에 사용하지 않는다. A₀ → B₀ → A₀의 FPS 크기는 2, '
        'A₀ → B₀ → A₁은 3이다. 후자에서 B₀ → A₁이 재등장해도 크기는 3으로 유지된다. '
        '따라서 거리 dₙ은 왕복 횟수나 편도 번역 수와 다르다.')
    paragraph(doc,'Sym_adj(t)는 xₜ₋₁과 xₜ, Sym_origin(t)는 x₀와 xₜ의 유사도다. '
        '모든 생성 단계가 기능 검사를 통과하고, Sym_adj ≥ 0.85 및 Dⱼ ≤ 10을 처음 만족한 복원 단계에서 '
        '종료하여 dₙ = Dⱼ, Sym_final = Sym_origin을 기록한다. 이는 정확한 동일성이나 영구 고정점의 증명이 아닌 '
        '운영적 근사 수렴이다. 기능 실패를 먼저 판정하고, FPS=11이면 해당 코드의 평가 후 추가 번역을 중지한다. '
        '중복·진동은 별도 편도 20회 예산으로 제한하며 추가 확인 번역은 수행하지 않는다.')
    paragraph(doc,'P_sf는 평가 가능한 관측 중 기능과 수렴의 공동 성공 비율이다. 종료까지 모든 생성 코드가 '
        '기능 검사를 통과한 비율 P_functional도 별도로 집계한다. 실패·상한 도달의 성공 조건부 dₙ과 Sym_final은 '
        'NA로 두고 실제 종료 FPS 크기를 보존한다. 인프라·유사도 오류와 미완료는 평가 실패와 구분한다.')
    paragraph(doc,'수치·문자열·리스트·동적 계획·탐색/관계의 5영역에서 번역 전에 각 3문제를 선정했다. '
        '세 언어 기준 코드 45개는 동일한 문제당 13개 평가 입력, 총 585개 검사를 통과했다. '
        '공개 예제는 프롬프트에 제공하고 평가 입력은 제공하지 않았다. 예비 3문제×6방향=18건은 본 실험과 분리했다. '
        '본 실험은 15문제×6방향×1모델×1반복=90건이며, 예비 문제를 제외한 12문제 분석도 수행했다.')
    paragraph(doc,'LM Studio의 Qwen2.5-Coder 7B Instruct(Q4_K_M), temperature=0, top_p=1, seed=20260918, '
        '출력 상한 4,096토큰을 사용했다. 도구는 GCC 14.2.0, GHC 9.10.3, SWI-Prolog 10.0.2, '
        'JPlag 6.3.0과 OpenJDK 26이다. 세 언어 모두 text 모드·최소 매칭 9토큰·averageSimilarity를 쓰고 '
        '입출력 래퍼를 포함했다. 생성 오류는 수정하거나 성공할 때까지 재생성하지 않았다.')
    paragraph(doc,'3. 실험 결과와 해석',11,True,4,True)
    paragraph(doc,'표 1. 방향별 성공률과 성공 조건부 거리·유사도',9,True,3,True)
    table=doc.add_table(rows=1, cols=4);table.autofit=False
    widths=[1.05,2.85,2.1,2.45]
    headers=['경로','P_sf [95% CI]','dₙ [Q1,Q3]','Sym_final']
    for c,w,t in zip(table.rows[0].cells,widths,headers):c.width=Cm(w);c.text=t
    for c,w in zip(table.columns,widths):c.width=Cm(w)
    for route,s in summary['routes'].items():
        a,b=route.split('-via-');ci=s['wilson95'];d=s['d_n']
        vals=[LABEL[a]+'→'+LABEL[b],f'{s["success"]}/15 ({s["p_sf"]*100:.0f}%)\n[{ci[0]*100:.1f},{ci[1]*100:.1f}]',
              'NA' if not d['n'] else f'{d["median"]:g} [{d["q1"]:g},{d["q3"]:g}]',
              'NA' if not d['n'] else f'{s["sym_final"]["median"]:.3f}']
        for c,w,t in zip(table.add_row().cells,widths,vals):c.width=Cm(w);c.text=t
    for row in table.rows:
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1
                for run in p.runs:font(run,9)
    paragraph(doc,'C=C++, H=Haskell, P=Prolog. A→B는 A→B→A 왕복 경로이며 각 15건이다. '
        '구간은 Wilson 95%, dₙ은 중앙값 [사분위], Sym_final은 성공 조건부 중앙값이다. NA는 성공 표본 없음이다.',9,False,5)
    failures=', '.join(f'{name} {counts[status]}건' for status,name in [('compile_error','컴파일/로딩 오류'),('runtime_error','실행 오류'),
                    ('wrong_answer','오답'),('timeout','시간 초과'),('compile_timeout','컴파일 시간 초과'),('max_distance_reached','FPS 상한'),
                    ('translation_budget_reached','번역 예산 상한'),('format_error','형식 오류')] if counts[status])
    pf_equal=all(s['p_sf']==s['p_functional'] for s in summary['routes'].values())
    paragraph(doc,f'90건 중 성공 {len(ok)}건, 평가 실패 {90-len(ok)}건이며 미평가·미완료는 0건이다. 실패는 {failures}이었다. '+
        ('기능을 통과한 미수렴 관측이 없어 P_functional은 각 경로의 P_sf와 같았다. ' if pf_equal else
         '기능 통과와 수렴 성공이 다른 관측은 별도로 집계하였다. ')+
        '성공 0건인 경로의 표본 bootstrap 구간은 [0,0]으로 퇴화하므로 Wilson 구간을 병기했다.')
    if ok:
        case=sorted(ok,key=lambda r:(r['problem_id'],r['route']))[0]
        a,b=case['route'].split('-via-')
        paragraph(doc,f'사전 규칙으로 선택한 성공 사례 {case["problem_id"]}({LABEL[a]}→{LABEL[b]})는 '
            f'{case["completed_roundtrips"]}회 왕복 후 dₙ={case["d_n"]}, Sym_adj={case["sym_adj"]:.3f}, '
            f'Sym_final={case["sym_final"]:.3f}였다. 인접 복원 코드의 안정화가 원본 표현의 복원을 뜻하지 않음을 보여준다. '
            '실패의 작은 종료 FPS는 가까운 언어의 근거가 아니다.')
    paragraph(doc,'이 팰린드롬 사례의 원본 C++는 양끝 문자를 인덱스로 비교했고, 안정화한 코드는 '
        'std::reverse로 만든 문자열과 원문을 비교했다. 두 방식 모두 13개 입력을 통과했지만 '
        'text의 최소 9토큰 일치 기준에서 원본 유사도는 0이었다. 마지막 두 편도의 코드가 이미 관측한 상태로 '
        '재등장하여, 6회 편도 번역과 출발 코드를 포함해도 서로 다른 상태는 5개였다.')
    paragraph(doc,'같은 문제의 C→P 경로는 첫 Prolog 코드가 존재하지 않는 read_line_to_string/1을 '
        '호출하여 입력 a에서 실행 오류를 냈다. 관측 FPS는 2지만 성공 거리 dₙ은 NA다. '
        'Prolog를 포함한 네 경로의 성공 0/15는 이 모델·프롬프트·입출력 계약에서의 관측이며, '
        '모집단 성공률 0이나 언어 자체의 거리 순위를 뜻하지 않는다.')
    paragraph(doc,'C→H와 H→C의 대응 성공률 차이(C→H에서 H→C를 뺌)는 -6.7%p이고 '
        'Bonferroni 구간은 [-33.3, 20.0]%p였다. 두 방향에서 모두 성공한 문제는 3개이며 '
        '거리의 평균 대응 차이는 0, 구간은 [-2,3]이었다. 이 표본만으로 방향 차이를 확정하기 어렵다.')
    sensitivity=summary['excluding_pilot_problems']['routes']
    reduced=sum(s['success'] for s in sensitivity.values())
    paragraph(doc,f'예비 사용 문제를 제외한 72건에서는 {reduced}건이 성공했다. '
        '반대 방향 3쌍은 같은 문제를 대응시켜 비교하고, 성공률은 공통 평가 문제, 거리·유사도는 공통 성공 문제만 사용했다. '
        '문제 단위 bootstrap 2,000회와 지표별 3비교 Bonferroni 구간을 적용했다. 공통 성공 표본이 없으면 '
        '거리 차이를 산출하지 않았다. 단일 반복으로 실행 간 변동성을 추정하지 않는다.')
    paragraph(doc,'사전 진단에서 동일 코드의 반복 측정은 1.0이었지만 곱셈을 덧셈으로 바꾼 코드도 1.0이었다. '
        'Prolog의 공통 입출력 부분만 남긴 점수도 0.852였다. JPlag text[2]는 표면 텍스트 척도이므로 '
        '이 점수나 임계값을 의미 동등성으로 해석할 수 없다. 기능 검사는 그 한계를 보완하지만 '
        '유한한 13개 입력만으로 모든 입력의 정확성을 보장하지 않는다.')
    # Place the summary figure and conclusion in the second page's right column.
    p=paragraph(doc,'',after=0);p.add_run().add_break(WD_BREAK.COLUMN)
    figure=root/'abstract_figure.png'
    make_figure(summary,records,figure)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2)
    p.add_run().add_picture(str(figure),width=Cm(8.3))
    paragraph(doc,'Fig. 1. 방향별 P_sf와 Wilson 95% 구간(위), 성공한 9건의 FPS 거리(아래). '
        '성공 사례가 없는 경로는 아래 그림에 표시하지 않았다. 점의 작은 세로 이동은 겹침을 줄이기 위한 표시이며 값의 변동이 아니다.',9,True,6)
    paragraph(doc,'4. 결론',11,True,4,True)
    paragraph(doc,'RTT의 성공률과 FPS 상태 수를 함께 제시하면 번역 실패와 성공한 표현 안정화를 구분할 수 있다. '
        '다만 dₙ은 성공 사례에 조건부여서 선택 편향을 가지며, 원본 표현과 입출력 래퍼·모델·도구·임계값에 의존한다. '
        '반대 방향은 출발 코드가 달라 대칭성을 전제할 수 없고 삼각부등식도 입증되지 않았다. '
        '따라서 수학적 거리함수나 패러다임의 인과 효과를 주장하지 않는다. 작은 공통 표본과 한 모델의 결과를 '
        '다른 모델·실행 반복으로 확장 검증하는 것이 후속 과제다.')
    paragraph(doc,'참고문헌',11,True,3,True)
    paragraph(doc,'[1] 조찬우 외, "왕복 번역과 고정점을 이용한 언어 거리 계산", 2025 한국소프트웨어종합학술대회 논문집, pp. 1467-1469, 2025.',9,False,3)
    paragraph(doc,'[2] JPlag, "JPlag: Source Code Plagiarism Detection", version 6.3.0, https://github.com/jplag/JPlag (접근: 2026-09-18).',9,False,0)
    for attr in ['author','last_modified_by','comments','keywords','subject']:setattr(doc.core_properties,attr,'')
    doc.core_properties.title='왕복 번역과 FPS 거리 실험'
    doc.save(destination)
    write_json(root/'abstract_provenance.json',{'at':timestamp(),'summary_sha256':digest((root/'summary.json').read_bytes()),
              'document_sha256':digest(destination.read_bytes()),'builder_sha256':digest(Path(__file__).read_bytes())})
    print(destination)

def make_figure(summary,records,path):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import numpy as np
    plt.rcParams.update({'font.size':8.5})
    fig,axs=plt.subplots(2,1,figsize=(3.25,3.8),gridspec_kw={'height_ratios':[2,1.1]})
    routes=list(summary['routes']);labels=[]
    for i,route in enumerate(routes):
        a,b=route.split('-via-');labels.append(LABEL[a]+'-'+LABEL[b]);r=summary['routes'][route]
        p=r['p_sf'];ci=r['wilson95']
        axs[0].errorbar(p,i,xerr=[[max(0,p-ci[0])],[max(0,ci[1]-p)]],fmt='o',color='#235887',capsize=3)
        axs[0].text(.63,i,f'{r["success"]}/15',va='center',fontsize=8)
    axs[0].set_yticks(range(6),labels);axs[0].invert_yaxis();axs[0].set_xlim(-.025,.8)
    axs[0].set_xlabel('P_sf (Wilson 95% interval)');axs[0].grid(axis='x',alpha=.2)
    for i,route in enumerate(['cpp-via-haskell','haskell-via-cpp']):
        vals=[r['d_n'] for r in records if r['route']==route and r['success']]
        axs[1].scatter(vals,i+np.linspace(-.09,.09,len(vals)),s=26,color='#235887',alpha=.8)
    axs[1].set_yticks([0,1],['C-H','H-C']);axs[1].set_ylim(-.5,1.5);axs[1].invert_yaxis()
    axs[1].set_xticks(range(2,8));axs[1].set_xlabel('Successful FPS size d_n');axs[1].grid(axis='x',alpha=.2)
    for ax in axs:
        ax.spines[['top','right']].set_visible(False)
    fig.tight_layout(h_pad=1.1);fig.savefig(path,dpi=300,bbox_inches='tight');plt.close(fig)

if __name__=='__main__':main()
